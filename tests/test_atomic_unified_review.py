"""Review regressions: recorded material, real Worker effects and intercepted HTTP."""
from copy import deepcopy
from dataclasses import asdict
from pathlib import Path
from types import SimpleNamespace
import json
import os
import zipfile

import pytest
import yaml

from atomic_skillgraph.core.errors import BudgetExhausted
from atomic_skillgraph.empirical.budget_governor import BudgetGovernor, input_token_bound
from atomic_skillgraph.empirical.contracts import PublicTask, digest, object_schema
from atomic_skillgraph.empirical.local_validation import (
    artifact_identity, artifact_record, canonical_binding, evidence_from_trace, input_sources,
    resolve_evidence, validate_on_source)
from atomic_skillgraph.empirical.model_view import builder_example, expand_material, learning_case
from atomic_skillgraph.empirical.prompts import BUILD, BUILDER_PROMPT, LEARNING, LEARNER_PROMPT
from test_empirical import config_for, worker
from test_atomic_unified import setup, SOURCE, SCHEMA, OUT
from test_cf2_contracts import http, response


def record(name,value):
    root=os.environ.get('ATOMIC_REVIEW_EVIDENCE_DIR')
    if root:
        path=Path(root);path.mkdir(parents=True,exist_ok=True)
        (path/(name+'.json')).write_text(json.dumps(value,ensure_ascii=False,indent=2))


def recorded(benchmark,task_id):
    root=os.environ.get('ATOMIC_REVIEW_ROOT')
    if not root:pytest.skip('Recorded review package not configured')
    paths=Path(root).joinpath('finite_batch',benchmark,'train','Train').glob('*/trace.json')
    trace=next(json.loads(p.read_text()) for p in paths if json.loads(p.read_text())['task']['task_id']==task_id)
    return trace


@pytest.mark.parametrize('benchmark,task_id',[
    ('officeqa','officeqa:UID0170'),('officeqa','officeqa:UID0221'),
    ('spreadsheetbench','spreadsheet:44296'),('spreadsheetbench','spreadsheet:58032')])
def test_recorded_complete_payloads_fit_learning_pool(tmp_path,monkeypatch,benchmark,task_id):
    from atomic_skillgraph.empirical.system import EmpiricalSystem
    from atomic_skillgraph.experiments.run_formal import resolved_config
    from atomic_skillgraph.harness.registry import create_simple_harness
    trace=recorded(benchmark,task_id);task=PublicTask(**trace['task'])
    base=config_for(tmp_path/'bank')
    base['budget']={'token_limit':200000,'request_limit':50,'finish_reserve':0,
        'train_task_tokens':200000,'train_solve_tokens':120000,'train_learning_tokens':80000,'eval_task_tokens':96000}
    profiles=json.loads(Path('benchmark_profiles.json').read_text())['profiles']
    cfg=resolved_config(base,profiles['spreadsheet' if benchmark=='spreadsheetbench' else benchmark],benchmark,42,
        'train',tmp_path,Path(os.environ['CF4_DATASETS']),Path('data/main_experiment_v1'),
        Path('/mnt/d/T3S_exp/SkillCompiler_resources_20261003/raw/officeqa/treasury_bulletins_parsed/transformed'))
    adapter=create_simple_harness(cfg)
    s=EmpiricalSystem(cfg,harness=adapter)
    sent=http(monkeypatch,[response([{'decision':'no_change'}],name='submit_learning'),response([{'source':SOURCE}],name='submit_program')])
    try:
        s._learning_start=0;s.phase='train'
        s.request_attribution={'parent_scope':task_id,'parent_token_limit':200000,'train_parent':True}
        result=s.learner.learn(task,trace)
        assert result['decision']=='no_change',result
        payload=sent[0]; assert input_token_bound(payload)<=16384
        material=expand_material(json.loads(payload['messages'][1]['content']))
        assert material['experience']['case_source']['physical_case_id']==task.physical_key
        assert all(h['case_source']['physical_case_id']!=task.physical_key for h in material['completed_train_cases'])
        experience=s.learner._experience(task,trace)
        assert digest(experience)==digest(s.bank.train_cases()[0]['experience'])
        # Select one existing complete operation, never a fabricated summary.
        e=next(e for e in experience['local_evidence'] if e['id']==material['experience']['local_evidence'][0]['local_evidence_ref'])
        sources=input_sources(e)
        binding={'case_id':task.physical_key,'binding_hash':'offline-probe','inputs':{},'reference_fields':{}}
        from atomic_skillgraph.empirical.program_worker import program_permission_view
        from atomic_skillgraph.empirical.task_context import TaskContext
        permissions=program_permission_view(adapter.tool_definitions(),[*adapter.available_tools(),TaskContext(cfg['runtime']).tool()],
            tool_surface=adapter.capabilities.tool_surface,workspace=getattr(adapter,'workspace',None),environment=cfg['program_environment'])
        build=s.learner.builder_material({'goal':'Replay selected actual operation','input_schema':object_schema(),
            'output_schema':object_schema()},[binding],[builder_example(binding,e)],permissions)
        s.agent('tool_builder',BUILDER_PROMPT,build,'submit_program',BUILD)
        assert input_token_bound(sent[1])<=16384
        actual=expand_material(json.loads(sent[1]['messages'][1]['content']))
        assert actual['examples'][0]['operation']==e['operation']
        assert 'prefix' not in actual['examples'][0]
        assert sources and material['experience']['local_evidence']
        ledger=list(s.budget_governor.state['attempts'].values())
        assert len(ledger)==2 and all(r['context']['budget_pool']=='learning' for r in ledger)
        assert sum(r['reserved_tokens'] for r in ledger)<=80000
        record(task_id.replace(':','_'),{'extractor_bytes':input_token_bound(payload),'builder_bytes':input_token_bound(sent[1]),
            'source_evidence_id':e['id'],'source_operation_hash':digest(e['operation']),
            'reserved_learning_tokens':sum(r['reserved_tokens'] for r in ledger),'result':result})
    finally:s.close()


def test_old_sheet_proposal_errors_aggregate_and_source_is_not_replaced(tmp_path):
    trace=recorded('spreadsheetbench','spreadsheet:382-10')
    s,t=setup(tmp_path)
    from atomic_skillgraph.harness.benchmarks import SpreadsheetAdapter
    s.adapter=SpreadsheetAdapter({},s.config)
    try:
        task=PublicTask(**trace['task']);experience=s.learner._experience(task,trace)
        proposal=next(r['response']['tool_calls'][0]['arguments'] for r in trace['requests']
            if r['stage']=='extractor' and r.get('response',{}).get('tool_calls'))
        old=digest(proposal)
        with pytest.raises(ValueError) as exc:s.learner.validate_learning_proposal(proposal,{task.physical_key:(task,experience)})
        codes={e['code'] for row in exc.value.feedback for e in row['errors']}
        assert {'host_authority_field','unbound_input','missing_output_reference','invalid_output_reference','no_publication'}<=codes
        assert digest(proposal)==old
        record('old_sheet_proposal',{'errors':exc.value.feedback,'original_proposal_unchanged':True})
    finally:s.close()


def test_input_provenance_and_host_binding(tmp_path):
    from atomic_skillgraph.empirical.local_validation import source_literals
    e={'id':'local:write','prefix':[{'name':'read','arguments':{}}],'source_trace_sha256':'source',
       'environment_identity':{},'public_task':{'goal':'use column A','inputs':{'n':3}},
       'operation':{'source':"sheet = 'Sheet1'\nif False:\n    value = 'invented'\n"},
       'reference':{'value':6}}
    sheet=next(r for r in source_literals(e['operation']['source']) if r['value']=='Sheet1')
    binding={'case_id':'physical','local_evidence_ref':e['id'],'inputs':{'n':3},
        'input_refs':{'n':{'kind':'public_json','path':['inputs','n']}},'reference_fields':{'value':['value']}}
    locked=canonical_binding(binding,{'local_evidence':[e]},{'input_schema':SCHEMA,'output_schema':OUT})
    assert locked['prefix']==e['prefix'] and locked['start_mode']=='prefix_replay' and locked['binding_hash']
    assert all(r['value']!='invented' for r in source_literals(e['operation']['source']))
    from atomic_skillgraph.empirical.local_validation import bound_input
    assert bound_input(e,sheet)=='Sheet1'
    assert bound_input(e,{'kind':'public_span','path':['goal'],'start':11,'end':12})=='A'
    for bad in ({'kind':'public_json','path':['inputs','gold']},{'kind':'public_span','path':['goal'],'start':99,'end':100}):
        with pytest.raises(ValueError):bound_input(e,bad)
    altered=deepcopy(locked);altered['prefix']=[]
    with pytest.raises(ValueError):resolve_evidence(altered,{'local_evidence':[e]},{'input_schema':SCHEMA,'output_schema':OUT})


def test_structural_repair_actual_purpose_payload_and_recovery(tmp_path,monkeypatch):
    from atomic_skillgraph.empirical.checkpoint import TaskCheckpoint
    s,t=setup(tmp_path);s.checkpoint=TaskCheckpoint(tmp_path/'checkpoint');s.phase='train';s._learning_start=0
    invalid={'decision':'reuse_existing','existing_skill_id':'nonexistent'}
    sent=http(monkeypatch,[response([invalid],name='submit_learning'),response([{'decision':'no_change'}],name='submit_learning')])
    try:
        answer=s.agent('extractor',LEARNER_PROMPT,{'experience':{'task':{'goal':'no recomputation'}}},'submit_learning',LEARNING,
            repair_limit=1,validator=s.learner.validate_learning_proposal,owner_state_version='repair-probe')
        assert answer['decision']=='no_change' and len(sent)==2
        assert sent[0]['thinking']['type']=='enabled'
        assert sent[1]['thinking']['type']=='disabled' and sent[1]['max_tokens']==2048
        assert sent[1]['tool_choice']['function']['name']=='submit_learning'
        assert input_token_bound(sent[1])<=8192
        assert [r['purpose'] for r in s.requests]==['extractor','learning_structure_repair']
        assert len({r['id'] for r in s.requests})==2
        assert s.agent('extractor',LEARNER_PROMPT,{'experience':{'task':{'goal':'no recomputation'}}},'submit_learning',LEARNING,
            repair_limit=1,validator=s.learner.validate_learning_proposal,owner_state_version='repair-probe')==answer
        assert len(sent)==2
        record('structural_repair',{'payloads':sent,'requests':s.requests})
    finally:s.close()


def test_recorded_sheet_proposal_one_short_repair(tmp_path,monkeypatch):
    from atomic_skillgraph.harness.benchmarks import SpreadsheetAdapter
    trace=recorded('spreadsheetbench','spreadsheet:382-10');task=PublicTask(**trace['task'])
    proposal=next(r['response']['tool_calls'][0]['arguments'] for r in trace['requests']
        if r['stage']=='extractor' and r.get('response',{}).get('tool_calls'))
    s,t=setup(tmp_path);s.adapter=SpreadsheetAdapter({},s.config)
    s._learning_start=0
    sent=http(monkeypatch,[response([proposal],name='submit_learning'),response([{'decision':'no_change'}],name='submit_learning')])
    try:
        experience=s.learner._experience(task,trace)
        result=s.agent('extractor',LEARNER_PROMPT,{'experience':learning_case(experience,task.physical_key)},
            'submit_learning',LEARNING,repair_limit=1,
            validator=lambda p:s.learner.validate_learning_proposal(p,{task.physical_key:(task,experience)}))
        assert result['decision']=='no_change' and len(sent)==2 and input_token_bound(sent[1])<=8192
        assert sent[1]['thinking']['type']=='disabled' and sent[1]['tool_choice']['function']['name']=='submit_learning'
        record('recorded_sheet_structure_repair',{'payloads':sent})
    finally:s.close()


def test_irrecoverable_truncation_and_repair_reasoning_only_reject(tmp_path,monkeypatch):
    def replay(trace, repair):
        request=next(r for r in trace['requests'] if r['stage']=='extractor' and r['repair']==repair)
        attempt=request['http_attempts'][-1]
        value=deepcopy(attempt['public_response']);value['usage']=deepcopy(attempt['raw_usage'])
        # Archive redacts private reasoning; restore only its empty transport envelope.
        value['choices'][0]['message']['reasoning_content']=''
        return value
    s,t=setup(tmp_path)
    sent=http(monkeypatch,[replay(recorded('officeqa','officeqa:UID0208'),0),
        response([{'decision':'reuse_existing','existing_skill_id':'missing'}],name='submit_learning'),
        replay(recorded('spreadsheetbench','spreadsheet:382-10'),1)])
    try:
        with pytest.raises(ValueError):s.agent('extractor',LEARNER_PROMPT,{},'submit_learning',LEARNING,repair_limit=1)
        assert len(sent)==1
        with pytest.raises(ValueError):s.agent('extractor',LEARNER_PROMPT,{},'submit_learning',LEARNING,repair_limit=1,
            validator=s.learner.validate_learning_proposal)
        assert len(sent)==3 and sent[-1]['max_tokens']==2048 and sent[-1]['thinking']['type']=='disabled'
        assert not s.bank.all('program')
    finally:s.close()


def test_final_payload_too_large_refused_without_http(tmp_path,monkeypatch):
    s,t=setup(tmp_path);sent=http(monkeypatch,[])
    try:
        with pytest.raises(ValueError,match='material_too_large'):
            s.agent('tool_builder',BUILDER_PROMPT,{'source':'x'*20000},'submit_program',BUILD)
        assert sent==[] and 'material_too_large' in s.requests[-1]['error']
    finally:s.close()


def test_train_pools_finish_learning_retry_and_identity(tmp_path):
    path=tmp_path/'budget.json';gov=BudgetGovernor(path,token_limit=200000,request_limit=20,finish_reserve=0)
    ctx={'parent_scope':'task','parent_token_limit':200000,'budget_pool':'solve','budget_pool_limit':120000,'parent_finish_reserve':5000}
    def charge(identity,total,context):
        gov.admit(identity,{'max_tokens':1},context)
        gov.complete({'request_id':identity,'outcome':'completed','raw_usage':{'total_tokens':total}})
    charge('runtime',114000,ctx)
    with pytest.raises(BudgetExhausted):gov.admit('runtime-refused',{'max_tokens':1100},ctx)
    finish={**ctx,'purpose':'finish_only','parent_finish_reserve':0}
    charge('finish',2000,finish)
    learn={**ctx,'budget_pool':'learning','budget_pool_limit':80000,'parent_finish_reserve':0,'stage':'runtime','phase':'trial'}
    charge('trial',65000,learn)
    gov.admit('retry',{'max_tokens':10000},learn)
    with pytest.raises(BudgetExhausted):gov.admit('over-learning',{'max_tokens':5000},learn)
    restored=BudgetGovernor(path,token_limit=200000,request_limit=20,finish_reserve=0)
    assert restored.state['unknown_billing']
    with pytest.raises(BudgetExhausted):restored.admit('unknown',{'max_tokens':1},finish)
    with pytest.raises(ValueError,match='identity changed'):BudgetGovernor(path,token_limit=200001,request_limit=20,finish_reserve=0)


def test_actual_solve_reservation_preserves_finish_and_learning_pool(tmp_path,monkeypatch):
    from atomic_skillgraph.empirical.answer_executor import STEP
    from atomic_skillgraph.empirical.prompts import finish_text_prompt
    s,t=setup(tmp_path);s.adapter.reset(t)
    s.config['budget']={'token_limit':200000,'request_limit':20,'finish_reserve':0,
        'train_task_tokens':200000,'train_solve_tokens':120000,'train_learning_tokens':80000,'eval_task_tokens':96000}
    s.budget_governor=BudgetGovernor(tmp_path/'budget.json',token_limit=200000,request_limit=20,finish_reserve=0)
    attribution={'parent_scope':t.physical_key,'parent_token_limit':200000,'train_parent':True}
    s.request_attribution=attribution;s.phase='train'
    gov=s.budget_governor
    gov.admit('prior',{'max_tokens':1},{**attribution,'budget_pool':'solve','budget_pool_limit':120000})
    gov.complete({'request_id':'prior','outcome':'completed','raw_usage':{'total_tokens':110000}})
    sent=http(monkeypatch,[response(content='6',finish='stop'),response([{'decision':'no_change'}],name='submit_learning')])
    try:
        material={'goal':t.goal,'inputs':t.inputs,'local_results':[],'local_reads':[]}
        with pytest.raises(BudgetExhausted):s.agent('runtime','Answer',material,'answer_step',STEP,finish_request_material=material)
        assert not sent
        assert s.agent('runtime',finish_text_prompt(s.adapter.answer_contract()),material,'finish_answer',None)=='6'
        s._learning_start=0
        assert s.agent('extractor','Learn',{},'submit_learning',LEARNING)['decision']=='no_change'
        ledger=list(gov.state['attempts'].values())
        assert [r['context']['budget_pool'] for r in ledger]==['solve','solve','learning']
        assert ledger[1]['context']['parent_finish_reserve']==0
        assert ledger[2]['context']['budget_pool_limit']==80000
    finally:s.close()


def test_office_runtime_reservation_uses_finish_evidence_contract(tmp_path,monkeypatch):
    from test_cf2_contracts import office
    s=office(tmp_path)
    s.config['budget']={'token_limit':200000,'request_limit':20,'finish_reserve':0,
        'train_task_tokens':200000,'train_solve_tokens':120000,'train_learning_tokens':80000,'eval_task_tokens':96000}
    s.budget_governor=BudgetGovernor(tmp_path/'budget.json',token_limit=200000,request_limit=20,finish_reserve=0)
    sent=http(monkeypatch,[response([{'action':'finish','answer':'private'}])])
    try:
        trace=s.run_task(s.adapter.task,learn=False)
        assert trace['score']['hard'] and len(sent)==1
        assert len(trace['requests'])==1 and not trace['requests'][0].get('error')
        admission=next(iter(s.budget_governor.state['attempts'].values()))['context']
        assert admission['budget_pool']=='solve' and admission['budget_pool_limit']==120000
        assert admission['parent_finish_reserve']>2048
    finally:s.close()


def test_xlsx_effect_time_and_semantic_changes(tmp_path):
    from openpyxl import Workbook, load_workbook
    from datetime import datetime
    original=tmp_path/'original.xlsx';wb=Workbook();wb.active['A2']='clear';wb.save(original)
    def save(name,change=None,time='2026-10-10T00:00:00Z'):
        p=tmp_path/name;book=load_workbook(original);book.active['A2']=None
        if change:change(book)
        book.save(p)
        # Deterministic different save instants, no sleep or model request.
        with zipfile.ZipFile(p) as archive:members={n:archive.read(n) for n in archive.namelist()}
        import xml.etree.ElementTree as ET
        core=ET.fromstring(members['docProps/core.xml']);core.find('{http://purl.org/dc/terms/}modified').text=time
        members['docProps/core.xml']=ET.tostring(core)
        with zipfile.ZipFile(p,'w') as archive:
            for n,data in members.items():archive.writestr(n,data)
        return p
    a=save('a.xlsx');b=save('b.xlsx',time='2026-10-10T00:00:02Z')
    assert artifact_identity(a)==artifact_identity(b)
    assert artifact_record(a)['raw_sha256']!=artifact_record(b)['raw_sha256']
    from openpyxl.styles import Font
    for name,change in [('value',lambda w:setattr(w.active['B2'],'value',7)),
                        ('formula',lambda w:setattr(w.active['B2'],'value','=1+2')),
                        ('style',lambda w:setattr(w.active['A2'],'font',Font(bold=True))),
                        ('sheet',lambda w:w.create_sheet('new')),
                        ('creator',lambda w:setattr(w.properties,'creator','another author'))]:
        assert artifact_identity(a)!=artifact_identity(save(name+'.xlsx',change))
    record('xlsx_identity',{'same_effect':artifact_record(a),'different_save':artifact_record(b)})


@pytest.mark.parametrize('benchmark',['searchqa','livemath','officeqa','spreadsheetbench'])
def test_formal_and_finite_resolve_same_method_and_static_budget(tmp_path,benchmark):
    from atomic_skillgraph.empirical.system import validate_config
    from atomic_skillgraph.experiments.run_formal import resolved_config, model_settings
    from atomic_skillgraph.experiments.run_atomic_unified_validation import LIMITS,PARENT_LIMITS
    from atomic_skillgraph.experiments.seed42_budget import allocation
    from atomic_skillgraph.experiments.canonical_manifest import verify
    spec=yaml.safe_load(Path('configs/atomic_unified_seed42.yaml').read_text())
    base=model_settings(yaml.safe_load(Path(spec['base_config']).read_text()),json.loads(Path(spec['model_lock']).read_text())['current_test'])
    profiles=json.loads(Path(spec['benchmark_profiles']).read_text())['profiles']
    frozen=json.loads(Path(spec['budget_allocation']).read_text())
    assert frozen==allocation(verify(spec['authority']),profiles,base['llm']['max_retries'])
    assert frozen['total_limits']['token_limit']==315232000
    formal=deepcopy(base);formal['budget']=frozen['cells'][benchmark]
    finite=deepcopy(base);finite['budget']={**LIMITS,**PARENT_LIMITS}
    profile=profiles['spreadsheet' if benchmark=='spreadsheetbench' else benchmark]
    args=(profile,benchmark,42,'train',tmp_path,Path('/public/datasets'),Path(spec['authority']),Path('/public/corpus'))
    a,b=(resolved_config(c,*args) for c in (formal,finite))
    aa,bb=deepcopy(a),deepcopy(b);aa.pop('budget');bb.pop('budget');assert aa==bb
    for cfg in (a,b):
        assert cfg['budget']['train_solve_tokens']+cfg['budget']['train_learning_tokens']==200000
        assert cfg['llm']['purpose_overrides']['learning_structure_repair']['max_completion_tokens']==2048
    if benchmark=='livemath':
        assert a['learning']['choice_guidance']['enabled'] and a['runtime']['choice_guidance']['enabled']
        assert a['llm']['purpose_overrides']['guidance_grounding']['max_completion_tokens']==1536
    else:assert not a['learning'].get('choice_guidance',{}).get('enabled')
    record('configuration_'+benchmark,{'formal':a,'finite':b,'equal_excluding_declared_budget':True})


def test_xlsx_publication_real_worker_with_host_binding(tmp_path,worker):
    from openpyxl import Workbook
    from atomic_skillgraph.empirical.system import EmpiricalSystem
    from atomic_skillgraph.harness.benchmarks import SpreadsheetAdapter
    from atomic_skillgraph.harness.simple_protocol import Broker
    from atomic_skillgraph.empirical.program_worker import ProgramWorker
    path=tmp_path/'input.xlsx';book=Workbook();book.active['A2']='clear';book.save(path)
    task=PublicTask('sheet','sheet-physical','Clear A2 in input workbook',{'files':[str(path)]})
    cfg=config_for(tmp_path/'bank');cfg['harness']={'adapter':'spreadsheet'}
    records={'sheet':{'public_files':{'input.xlsx':path}}}
    adapter=SpreadsheetAdapter(records,cfg)
    s=EmpiricalSystem(cfg,harness=adapter,adapter_factory=lambda:SpreadsheetAdapter(records,cfg))
    s.worker=ProgramWorker({**cfg['program_worker'],'wall_timeout_seconds':60})
    adapter.worker=s.worker
    try:
        adapter.reset(task)
        source="from openpyxl import load_workbook\nw=load_workbook(INPUT_PATH)\nw.active['A2']=None\nw.save(OUTPUT_PATH)"
        broker=Broker(adapter,30);broker.call('execute_python',{'source':source,'files':['case1_result.xlsx']})
        event=broker.events[-1];assert event['result']['accepted'],json.dumps(event)
        trace={'tools':broker.events,'execution':{}};e=evidence_from_trace(task,trace,s.config['program_environment'])[0]
        input_schema=object_schema({'path':{'type':'string'},'cell':{'type':'string'}},['path','cell'])
        output_schema=object_schema({'files':{'type':'array','items':{'type':'string'}},'output_file':{'type':'string'}},['files','output_file'])
        literal=next(r for r in input_sources(e) if r['value']=='A2')
        binding={'case_id':task.physical_key,'local_evidence_ref':e['id'],
            'inputs':{'path':str(path),'cell':'A2'},'input_refs':{'path':{'kind':'public_json','path':['inputs','files',0]},'cell':literal},
            'reference_fields':{'files':{'kind':'publication','field':'files'},'output_file':{'kind':'publication','field':'file','name':'case1_result.xlsx'}}}
        declaration={'input_schema':input_schema,'output_schema':output_schema,'result_role':'intermediate'}
        locked=canonical_binding(binding,{'local_evidence':[e]},declaration)
        source="def run(ctx, inputs):\n    from openpyxl import load_workbook\n    w=load_workbook('/workspace/inputs/input.xlsx')\n    w.active[inputs['cell']]=None\n    w.save('/workspace/case1_result.xlsx')\n    return {'status':'ok','outputs':{'files':['case1_result.xlsx'],'output_file':'case1_result.xlsx'}}"
        p=s.bank.put('program',{**declaration,'source':source,'entry':'run','allowed_tools':[], 'environment':s.config['program_environment']})
        checked=validate_on_source(s,p,locked,task,{'local_evidence':[e]},'real-xlsx')
        assert checked['validation']['passed'],checked
        assert s.bank.program_eligible(s.bank.get(p['id']))
        assert checked['validation']['file_effect_version']==artifact_record(path)['algorithm']
        record('xlsx_worker_publication',{'source_event':event,'canonical_binding':locked,'validation':checked})
    finally:s.close()
